"""Safely enrich existing local patients from the approved NIS-matched workbooks."""
import argparse
from datetime import datetime
from pathlib import Path
import sqlite3

from openpyxl import load_workbook


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("database", type=Path)
    parser.add_argument("template", type=Path)
    parser.add_argument("source", type=Path)
    parser.add_argument("--school-id", required=True, type=int)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    if not args.database.is_file():
        raise ValueError("Database does not exist")
    template = load_workbook(args.template, read_only=True, data_only=True)
    expected = {}
    for row in template["Template Import Siswa"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        key = str(row[0]).strip()
        if key in expected:
            raise ValueError("Duplicate template NIS")
        expected[key] = str(row[1]).strip()
    template.close()
    source = load_workbook(args.source, read_only=True, data_only=True)
    values = {}
    for row in source["Data terbaru SISWA SRMA 13 BEKA"].iter_rows(min_row=6, values_only=True):
        if not row[1] or not row[5]:
            continue
        key, nik = str(row[5]).strip(), row[6]
        if key in values or not isinstance(nik, str) or len(nik.strip()) != 16 or not nik.strip().isascii() or not nik.strip().isdigit():
            raise ValueError("Invalid or duplicate source identifier; no changes made")
        values[key] = nik.strip()
    source.close()
    if not expected or set(values) != set(expected) or len(set(values.values())) != len(values):
        raise ValueError("Workbook identifiers do not match uniquely")
    db = sqlite3.connect(args.database)
    try:
        db.execute("BEGIN IMMEDIATE")
        columns = [r[1] for r in db.execute("PRAGMA table_info(patients)")]
        before = db.execute("SELECT * FROM patients ORDER BY id").fetchall()
        records = [dict(zip(columns, row)) for row in before]
        scoped = {str(r["id"]): r for r in records if r["school_id"] == args.school_id}
        if not scoped or not set(scoped).issubset(expected):
            raise ValueError("School roster differs from template; no changes made")
        print(f"Template rows absent from database: {len(set(expected) - set(scoped))}; will not recreate them.")
        for key in scoped:
            name = expected[key]
            if scoped[key]["name"].strip().casefold() != name.casefold():
                raise ValueError("Database name differs from approved template")
            if scoped[key].get("nik") not in (None, "", values[key]):
                raise ValueError("Existing NIK conflicts with source")
        print(f"Validated {len(scoped)} existing students; no new students required.")
        if not args.apply:
            db.rollback()
            return
        backup_dir = args.database.resolve().parent / "backups"
        backup_dir.mkdir(exist_ok=True)
        backup = backup_dir / f"before_patient_nik_{datetime.now():%Y%m%d_%H%M%S_%f}.db"
        with sqlite3.connect(args.database) as reader, sqlite3.connect(backup) as destination:
            reader.backup(destination)
        if "nik" not in columns:
            db.execute("ALTER TABLE patients ADD COLUMN nik VARCHAR(16)")
        db.execute("CREATE INDEX IF NOT EXISTS ix_patients_nik ON patients (nik)")
        for key in scoped:
            nik = values[key]
            db.execute("UPDATE patients SET nik=? WHERE id=? AND school_id=?", (nik, key, args.school_id))
        preserved = [c for c in columns if c != "nik"]
        selected = ", ".join('"' + c + '"' for c in preserved)
        after = db.execute(f"SELECT {selected} FROM patients ORDER BY id").fetchall()
        original = [tuple(r[c] for c in preserved) for r in records]
        if after != original:
            raise ValueError("Unexpected non-NIK changes; transaction rolled back")
        db.commit()
        print(f"Updated NIK for {len(scoped)} students. All other patient fields unchanged.")
        print(f"Backup: {backup}")
    finally:
        db.close()


if __name__ == "__main__":
    main()
